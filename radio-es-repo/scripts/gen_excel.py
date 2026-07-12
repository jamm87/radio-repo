import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

data = json.load(open('consolidado.json'))

# ---- normalizar frecuencia a número (MHz) para poder ordenar/filtrar ----
def to_mhz(row):
    v = row['Frecuencia']; u = row['Unidad']
    if isinstance(v, (int, float)):
        f = float(v)
    else:
        s = str(v)
        m = re.match(r'^\s*([\d]{1,3}(?:[.,]\d+)?)', s)
        if not m: return None
        f = float(m.group(1).replace(',', '.'))
    if u == 'kHz': f = f/1000
    elif u == 'GHz': f = f*1000
    return round(f, 6)

for r in data:
    r['MHz'] = to_mhz(r)

# ordenar: primero por publicable (SI, DUDOSO, NO), luego por frecuencia
order = {'SI':0,'DUDOSO':1,'NO':2}
data.sort(key=lambda r:(order.get(r['Publicable'],3), r['MHz'] if r['MHz'] is not None else 9e9))

wb = Workbook()

# ================= HOJA 1: FRECUENCIAS =================
ws = wb.active
ws.title = "Frecuencias"

cols = ['Frecuencia','Unidad','MHz (norm.)','Modo','Descripción','Categoría','Zona','Fuente','Publicable','Motivo clasificación','Notas']
headfill = PatternFill('solid', fgColor='1C242C')
headfont = Font(name='Arial', bold=True, color='FFB454', size=10)
thin = Side(style='thin', color='D9D9D9')
border = Border(left=thin,right=thin,top=thin,bottom=thin)

for c,h in enumerate(cols,1):
    cell = ws.cell(1,c,h)
    cell.fill = headfill; cell.font = headfont
    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    cell.border = border

fills = {
    'SI':   PatternFill('solid', fgColor='E5F5E0'),   # verde claro
    'DUDOSO':PatternFill('solid', fgColor='FFF2CC'),  # ámbar claro
    'NO':   PatternFill('solid', fgColor='F8D7DA'),   # rojo claro
}
pubfont = {
    'SI':   Font(name='Arial', bold=True, color='1E7B34', size=10),
    'DUDOSO':Font(name='Arial', bold=True, color='9C6500', size=10),
    'NO':   Font(name='Arial', bold=True, color='B02A37', size=10),
}
base = Font(name='Arial', size=10)

for i,r in enumerate(data, start=2):
    vals = [r['Frecuencia'], r['Unidad'], r['MHz'], r['Modo'], r['Descripción'],
            r['Categoría'], r['Zona'], r['Fuente'], r['Publicable'],
            r['Motivo clasificación'], r['Notas']]
    pub = r['Publicable']
    for c,v in enumerate(vals,1):
        cell = ws.cell(i,c,v)
        cell.font = base
        cell.border = border
        cell.alignment = Alignment(vertical='top', wrap_text=(c in (5,10,11)))
        # colorear toda la fila suave según publicable
        cell.fill = fills[pub]
        if c == 9:  # columna Publicable en negrita de color
            cell.font = pubfont[pub]
            cell.alignment = Alignment(horizontal='center', vertical='center')
        if c == 3 and v is not None:  # MHz normalizada
            cell.number_format = '0.000###'

# anchos
widths = [16,8,12,8,46,20,16,30,12,30,40]
for c,w in enumerate(widths,1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(data)+1}"
ws.row_dimensions[1].height = 30

# ================= HOJA 2: RESUMEN =================
ws2 = wb.create_sheet("Resumen")
title = Font(name='Arial', bold=True, size=14, color='1C242C')
h2 = Font(name='Arial', bold=True, size=11, color='FFFFFF')
h2fill = PatternFill('solid', fgColor='1C242C')
normal = Font(name='Arial', size=10)
boldn = Font(name='Arial', bold=True, size=10)

ws2['A1'] = 'RADIO://es — Consolidado de frecuencias'
ws2['A1'].font = title
ws2['A2'] = 'Extracción completa de los 4 documentos aportados, clasificada por publicabilidad.'
ws2['A2'].font = Font(name='Arial', size=10, italic=True, color='555555')

# tabla de conteo por publicable
from collections import Counter
cpub = Counter(r['Publicable'] for r in data)
ws2['A4']='CLASIFICACIÓN'; ws2['B4']='Nº frecuencias'; ws2['C4']='Significado'
for c in ('A4','B4','C4'):
    ws2[c].font=h2; ws2[c].fill=h2fill
rows_sum = [
    ('SI (publicable)', cpub.get('SI',0), 'Banda pública o de uso común: aeronáutica AIP, radioafición, CB, PMR446, marítimo, broadcast, meteo.'),
    ('DUDOSO (revisar)', cpub.get('DUDOSO',0), 'Requiere criterio: emergencias sanitarias, protección civil, transporte, utilities, "uso desconocido". Publicar solo tras revisión caso a caso.'),
    ('NO (no publicable)', cpub.get('NO',0), 'Fuerzas de seguridad, defensa/militar, seguridad privada, redes cifradas e infraestructura crítica. NO publicar.'),
]
for i,(k,v,s) in enumerate(rows_sum, start=5):
    ws2.cell(i,1,k).font=boldn
    ws2.cell(i,1).fill=fills[k.split()[0]]
    ws2.cell(i,2,v).font=normal
    ws2.cell(i,3,s).font=normal
    ws2.cell(i,3).alignment=Alignment(wrap_text=True, vertical='top')
ws2.cell(8,1,'TOTAL').font=boldn
ws2.cell(8,2,f'=SUM(B5:B7)').font=boldn

# tabla por fuente
ws2['A10']='POR FUENTE'; ws2['B10']='Nº frecuencias'
for c in ('A10','B10'): ws2[c].font=h2; ws2[c].fill=h2fill
cfu = Counter(r['Fuente'] for r in data)
for i,(k,v) in enumerate(sorted(cfu.items(), key=lambda x:-x[1]), start=11):
    ws2.cell(i,1,k).font=normal
    ws2.cell(i,2,v).font=normal

# tabla por categoria
start_cat = 11+len(cfu)+2
ws2.cell(start_cat-1,1,'POR CATEGORÍA').font=h2
ws2.cell(start_cat-1,1).fill=h2fill
ws2.cell(start_cat-1,2,'Nº').font=h2
ws2.cell(start_cat-1,2).fill=h2fill
ccat = Counter(r['Categoría'] for r in data)
for i,(k,v) in enumerate(sorted(ccat.items(), key=lambda x:-x[1]), start=start_cat):
    ws2.cell(i,1,k).font=normal
    ws2.cell(i,2,v).font=normal

# nota metodológica
note_row = start_cat+len(ccat)+2
ws2.cell(note_row,1,'CRITERIO DE PUBLICABILIDAD').font=boldn
notes = [
 'La columna "Publicable" evita problemas legales al difundir el listado. Se basa en el criterio del proyecto:',
 '• Frecuencias publicadas oficialmente (AIP de ENAIRE) o de bandas de uso común (radioafición, CB, PMR446, LPD433, marítimo) → SÍ.',
 '• Canales tácticos de fuerzas de seguridad, defensa, seguridad privada, redes cifradas (TETRA/SIRDEE) e infraestructura crítica → NO.',
 '• Emergencias sanitarias, protección civil, transporte público, utilities y "uso desconocido" → DUDOSO (revisar caso a caso).',
 'El CNAF advierte de no monitorizar comunicaciones de seguridad sin autorización; su contenido no debe divulgarse.',
 'Las frecuencias aeronáuticas cambian con los ciclos AIRAC: verificar contra el AIP en vigor antes de publicar como fiables.',
]
for i,t in enumerate(notes):
    ws2.cell(note_row+1+i,1,t).font=Font(name='Arial', size=9, color='555555' if t.startswith('•')==False else '000000')
    ws2.merge_cells(start_row=note_row+1+i, start_column=1, end_row=note_row+1+i, end_column=8)

ws2.column_dimensions['A'].width=42
ws2.column_dimensions['B'].width=16
ws2.column_dimensions['C'].width=70

wb.save('/mnt/user-data/outputs/frecuencias_radio_es.xlsx')
print("Excel generado:", len(data), "filas")
